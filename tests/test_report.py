import json

import pytest
from conftest import FakeOpenAI
from docx import Document

from mva.cli import main
from mva.llm import ExecutiveSummary, OpenAIExecutiveSummaryWriter
from mva.orchestrator import Orchestrator
from mva.report import build_report
from mva.review import ReviewSession


@pytest.fixture
def credit_run(settings, credit_pkg):
    return Orchestrator(settings).validate(credit_pkg)[1]


@pytest.fixture
def market_run(settings, market_pkg):
    return Orchestrator(settings).validate(market_pkg)[1]


def _text(path):
    doc = Document(str(path))
    body = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    return doc, body, cells


def test_draft_report_before_review(market_run):
    path = build_report(market_run)
    assert path.name.endswith("_DRAFT.docx")
    doc, body, cells = _text(path)
    assert "DRAFT" in body
    assert "Pending" in cells
    assert len(doc.inline_shapes) == 2  # backtest chart and challenger chart


def test_final_report_after_sign_off(credit_run):
    s = ReviewSession(credit_run)
    ids = s.pending()
    for fid in ids:
        s.accept(fid, "Hannah Zou")
    s.reject(ids[0], "Hannah Zou", "Raised elsewhere.")
    s.edit(ids[1], "Hannah Zou", "Downgraded after discussion with owner.", severity="Low")
    s.sign_off("Hannah Zou", "All findings reviewed.")
    path = build_report(credit_run)
    doc, body, cells = _text(path)
    assert not path.name.endswith("_DRAFT.docx") and "DRAFT" not in body
    for heading in ("1. Executive summary", "2. Model and validation scope", "3. Findings summary",
                    "4. Detailed findings", "Appendix A. Findings rejected by the validator",
                    "Appendix B. Audit trail and turnaround"):
        assert heading in body
    assert "Unsatisfactory" in body
    assert "Raised elsewhere." in cells  # rejected finding is in Appendix A with its rationale
    assert "changed by reviewer from" in body
    assert "Hannah Zou" in cells
    findings_table = next(t for t in doc.tables if t.rows[0].cells[0].text == "ID" and len(t.columns) == 6)
    assert len(findings_table.rows) - 1 == len(ids) - 1  # every retained finding, rejected one excluded
    assert len(doc.inline_shapes) == 2


def test_llm_executive_summary_is_guarded(market_run):
    fake = FakeOpenAI(ExecutiveSummary(summary="LLM summary text.", key_finding_ids=["BT-001", "ZZ-999"]))
    path = build_report(market_run, exec_writer=OpenAIExecutiveSummaryWriter("m", client=fake))
    _, body, _ = _text(path)
    assert "LLM summary text." in body and "drafted by an LLM" in body
    payload = json.loads(fake.calls[0]["input"][1]["content"])
    assert {"model", "rating", "findings"} <= set(payload)


def test_llm_failure_falls_back_to_template(market_run):
    fake = FakeOpenAI(error=RuntimeError("timeout"))
    path = build_report(market_run, exec_writer=OpenAIExecutiveSummaryWriter("m", client=fake))
    _, body, _ = _text(path)
    assert "overall rating of" in body
    events = [json.loads(x) for x in (market_run / "audit.jsonl").read_text().splitlines()]
    report_events = [e for e in events if e["event"] == "report_generated"]
    assert report_events[-1]["narrative"] == "template" and "timeout" in report_events[-1]["error"]


def test_cli_report_and_turnaround(market_run, capsys):
    assert main(["report", str(market_run), "--no-llm", "--out", str(market_run / "r.docx")]) == 0
    assert (market_run / "r.docx").exists()
    assert main(["turnaround", str(market_run)]) == 0
    assert '"automated_seconds"' in capsys.readouterr().out


def test_cli_benchmark(market_pkg, tmp_path, capsys):
    assert main(["benchmark", str(market_pkg), "--repeat", "2", "--runs-dir", str(tmp_path / "b")]) == 0
    result = json.loads((tmp_path / "b" / "benchmark.json").read_text())
    assert result["repeat"] == 2
    assert "total" in result["median_seconds"][str(market_pkg)]
