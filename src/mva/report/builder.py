"""Builds the model validation report from a reviewed run.

Before sign-off the document is marked DRAFT and pending findings are shown as
such; after sign-off it carries the reviewer, the date and the overall rating.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from mva.llm import (
    ExecutiveSummaryWriter,
    TemplateExecutiveSummaryWriter,
    TemplateNarrativeWriter,
    write_executive_summary,
)
from mva.report import charts
from mva.review import RATING_MEANING, ReviewSession, ReviewStatus
from mva.schemas import SR117Pillar
from mva.turnaround import turnaround

SEVERITY_FILL = {  # status palette; the label is always printed, so colour never carries meaning alone
    "High": ("D03B3B", "FFFFFF"),
    "Medium": ("FAB219", "0B0B0B"),
    "Low": ("E1E0D9", "0B0B0B"),
    "Info": ("F0EFEC", "0B0B0B"),
}
INK_2 = RGBColor(0x52, 0x51, 0x4E)
TYPE_LABEL = {"credit_pd": "Credit risk — probability of default (PD)", "market_var": "Market risk — Value-at-Risk"}


# -- small docx helpers -------------------------------------------------------------
def _shade(cell, fill: str, text_color: str | None = None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)
    if text_color:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor.from_string(text_color)
                r.font.bold = True


def _table(doc, header: list[str], rows: list[list[str]], widths: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        t.autofit = False
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        _shade(cell, "F0EFEC")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
            if widths:
                cell.width = Inches(widths[i])
    if widths:
        for i, col in enumerate(t.columns):
            col.width = Inches(widths[i])
    return t


def _severity_cell(cell, severity: str) -> None:
    fill, ink = SEVERITY_FILL.get(severity, ("FFFFFF", "0B0B0B"))
    _shade(cell, fill, ink)


def _para(doc, text: str, *, italic=False, bold=False, size=None, color=None, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.italic, run.bold = italic, bold
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return p


def _labelled(doc, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(text)


def _fmt(v: Any) -> str:
    if isinstance(v, float):
        return f"{v:.4g}"
    if isinstance(v, list):
        shown = ", ".join(_fmt(x) for x in v[:3])
        return f"[{shown}{', …' if len(v) > 3 else ''}]"
    if isinstance(v, dict):
        return "{" + ", ".join(f"{k}: {_fmt(x)}" for k, x in list(v.items())[:5]) + (", …" if len(v) > 5 else "") + "}"
    return str(v)


def _evidence_lines(evidence: dict[str, Any], limit: int = 8) -> list[str]:
    items = list(evidence.items())
    lines = [f"{k}: {_fmt(v)}" for k, v in items[:limit]]
    if len(items) > limit:
        lines.append(f"(+{len(items) - limit} more items in results.json)")
    return lines


# -- report ------------------------------------------------------------------------------
def _exec_payload(session: ReviewSession, manifest: dict) -> dict:
    finals = session.final_findings()
    rating = session.rating()
    return {
        "model": {"name": session.run.model_name, "id": session.run.model_id, "type": session.run.model_type,
                  "purpose": manifest.get("documentation", {}).get("purpose", ""),
                  "intended_use": manifest.get("documentation", {}).get("intended_use", "")},
        "rating": rating.value,
        "rating_meaning": RATING_MEANING[rating],
        "findings": [
            {"finding_id": x.finding.finding_id, "severity": x.finding.severity.value, "title": x.finding.title,
             "area": x.finding.agent, "description": x.finding.description}
            for x in finals if x.review.status is not ReviewStatus.PENDING
        ],
        "area_summaries": {r.agent: r.summary for r in session.run.agent_reports if r.status == "completed"},
    }


def _charts(session: ReviewSession, workdir: Path) -> list[tuple[str, Path, str]]:
    """(section heading, image, caption) for each chart the run's data supports."""
    out = []
    reports = {r.agent: r for r in session.run.agent_reports if r.status == "completed"}
    try:
        if session.run.model_type == "credit_pd":
            bt, ch = reports.get("Backtesting"), reports.get("Challenger Benchmarking")
            if bt and "calibration" in bt.metrics:
                out.append(("Backtesting", charts.credit_calibration(bt.metrics["calibration"]["deciles"], workdir / "calibration.png"),
                            "Mean predicted PD and observed default rate by decile of predicted PD."))
            if ch and "models" in ch.metrics:
                out.append(("Challenger Benchmarking", charts.credit_challengers(ch.metrics["models"], workdir / "challengers.png"),
                            "Validation-sample AUC of the delivered model and the validator's challengers."))
        else:
            bt, ch = reports.get("Backtesting"), reports.get("Challenger Benchmarking")
            art = session.run_dir / "artifacts" / "BT_var_backtest.csv"
            if bt and art.exists():
                conf = bt.metrics.get("var_config", {}).get("confidence", 0.99)
                out.append(("Backtesting", charts.var_backtest(pd.read_csv(art), conf, workdir / "var_backtest.png"),
                            "Exceptions are days whose loss exceeded the VaR forecast made the previous day."))
            if ch and "models" in ch.metrics:
                out.append(("Challenger Benchmarking", charts.var_challengers(ch.metrics["models"], workdir / "var_challengers.png"),
                            "Exception counts over the same backtest days; dashed line = expected count."))
    except Exception as exc:  # a chart must never block the report
        out.append(("_error", Path(), f"Charts could not be produced: {exc}"))
    return out


def build_report(
    run_dir: str | Path,
    out_path: str | Path | None = None,
    exec_writer: ExecutiveSummaryWriter | None = None,
) -> Path:
    session = ReviewSession(run_dir)
    run = session.run
    run_dir = Path(run_dir)
    snapshot = json.loads((run_dir / "manifest.json").read_text()) if (run_dir / "manifest.json").exists() else {}
    manifest = snapshot.get("manifest", {})
    doc_meta = manifest.get("documentation", {})
    signed = session.state.sign_off
    out_path = Path(out_path) if out_path else run_dir / f"validation_report_{run.model_id}{'' if signed else '_DRAFT'}.docx"

    payload = _exec_payload(session, manifest)
    exec_summary, exec_source, exec_error = write_executive_summary(exec_writer or TemplateExecutiveSummaryWriter(), payload)
    session.audit.event("report_generated", narrative=exec_source, error=exec_error, signed=bool(signed))

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)
        header = section.header.paragraphs[0]
        header.text = f"Model Validation Report · {run.model_id}" + ("" if signed else " · DRAFT — not signed off")
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = INK_2

    # -- cover -----------------------------------------------------------------------
    doc.add_heading("Model Validation Report", 0)
    _para(doc, run.model_name, bold=True, size=14)
    if not signed:
        _para(doc, "DRAFT — findings pending human review; not for distribution", bold=True,
              color=RGBColor(0xD0, 0x3B, 0x3B))
    cover = [
        ["Model ID / version", f"{run.model_id} / {manifest.get('version', '—')}"],
        ["Model type", TYPE_LABEL.get(run.model_type, run.model_type)],
        ["Model owner", doc_meta.get("owner", "—")],
        ["Validation run", run.run_id],
        ["Automated analysis", f"{run.started_at:%Y-%m-%d %H:%M} UTC ({run.duration_seconds:.1f} s)"],
        ["Reviewed and signed off by", f"{signed.reviewer}, {signed.signed_at:%Y-%m-%d %H:%M} UTC" if signed else "Pending"],
        ["Overall rating", session.rating().value + ("" if signed else " (provisional)")],
        ["Framework", "Federal Reserve / OCC SR 11-7 — Supervisory Guidance on Model Risk Management"],
    ]
    _table(doc, ["Item", "Detail"], cover, [2.0, 4.6])

    # -- 1. executive summary ---------------------------------------------------------------
    doc.add_heading("1. Executive summary", 1)
    rating = session.rating()
    p = doc.add_paragraph()
    p.add_run("Overall rating: ").bold = True
    p.add_run(rating.value + ("" if signed else " (provisional, review in progress)")).bold = True
    _para(doc, RATING_MEANING[rating], italic=True)
    doc.add_paragraph(exec_summary.summary)
    _para(doc, "Executive summary drafted by " + ("an LLM from the reviewed findings" if exec_source == "llm"
          else "template from the reviewed findings") + "; ratings and severities are set by the validator.",
          italic=True, size=8, color=INK_2)

    finals_all = session.final_findings(include_rejected=True)
    counted = [x for x in finals_all if x.review.status is not ReviewStatus.REJECTED]
    sev_rows = []
    for s in ("High", "Medium", "Low"):
        n = sum(x.finding.severity.value == s for x in counted)
        due = dict(session.policy.remediation_days).get(s)
        sev_rows.append([s, str(n), f"{due} days" if due else "—"])
    t = _table(doc, ["Severity", "Findings", "Remediation due within"], sev_rows, [1.5, 1.2, 2.0])
    for i, row in enumerate(t.rows[1:]):
        _severity_cell(row.cells[0], sev_rows[i][0])
    prog = session.progress()
    _para(doc, f"Review status: {prog['accepted']} accepted, {prog['edited']} edited, {prog['rejected']} rejected, "
          f"{prog['pending']} pending.", size=8.5, color=INK_2)

    # -- 2. scope -------------------------------------------------------------------------------
    doc.add_heading("2. Model and validation scope", 1)
    _labelled(doc, "Purpose", doc_meta.get("purpose", "—"))
    _labelled(doc, "Intended use", doc_meta.get("intended_use", "—"))
    if doc_meta.get("assumptions"):
        _para(doc, "Documented assumptions tested:", bold=True)
        for a in doc_meta["assumptions"]:
            doc.add_paragraph(a, style="List Bullet")
    if doc_meta.get("limitations"):
        _para(doc, "Documented limitations:", bold=True)
        for a in doc_meta["limitations"]:
            doc.add_paragraph(a, style="List Bullet")
    _para(doc, "Validation activities", bold=True)
    pillar_by_agent: dict[str, set] = {}
    for f in run.findings:
        pillar_by_agent.setdefault(f.agent, set()).add(f.pillar.value)
    default_pillar = {
        "Data Quality": SR117Pillar.CONCEPTUAL_SOUNDNESS.value,
        "Assumptions": SR117Pillar.CONCEPTUAL_SOUNDNESS.value,
        "Backtesting": SR117Pillar.OUTCOMES_ANALYSIS.value,
        "Challenger Benchmarking": SR117Pillar.ONGOING_MONITORING.value,
        "Sensitivity Analysis": SR117Pillar.CONCEPTUAL_SOUNDNESS.value,
    }
    rows = []
    for r in run.agent_reports:
        pillars = ", ".join(sorted(pillar_by_agent.get(r.agent, {default_pillar.get(r.agent, "—")})))
        rows.append([r.agent, pillars, "Completed" if r.status == "completed" else f"FAILED: {r.error}",
                     str(len(r.findings))])
    _table(doc, ["Review area", "SR 11-7 element", "Status", "Findings proposed"], rows, [1.6, 2.6, 1.4, 1.0])
    files = snapshot.get("file_sha256", {})
    if files:
        _para(doc, "Package validated (SHA-256 of each file, first 12 characters): "
              + "; ".join(f"{k} {v[:12]}" for k, v in files.items()), size=8, color=INK_2)
    notes = next((r.metrics.get("data_treatment") for r in run.agent_reports if r.metrics.get("data_treatment")), None)
    if notes:
        _para(doc, "Data treatment applied before model-level tests: " + "; ".join(notes) + ".", size=8.5, color=INK_2)

    # -- 3. findings summary -------------------------------------------------------------------
    doc.add_heading("3. Findings summary", 1)
    rows = [[x.finding.finding_id, x.finding.severity.value, x.finding.title, x.finding.agent,
             x.review.status.value.capitalize(), x.remediation_due or "—"] for x in counted]
    if rows:
        t = _table(doc, ["ID", "Severity", "Finding", "Area", "Review", "Due"], rows, [0.6, 0.7, 2.6, 1.2, 0.7, 0.8])
        for i, row in enumerate(t.rows[1:]):
            _severity_cell(row.cells[1], rows[i][1])
    else:
        doc.add_paragraph("No findings.")

    # -- 4. detailed findings by area -------------------------------------------------------------
    doc.add_heading("4. Detailed findings", 1)
    chart_list = _charts(session, Path(tempfile.mkdtemp()))
    for err in [c for c in chart_list if c[0] == "_error"]:
        _para(doc, err[2], italic=True, size=8, color=INK_2)
    for r in run.agent_reports:
        doc.add_heading(r.agent, 2)
        if r.status != "completed":
            _para(doc, f"This review area did not complete: {r.error}", italic=True)
            continue
        reviewed = [x.finding for x in counted if x.finding.agent == r.agent]
        if r.narrative_source == "template":
            # Rebuild the deterministic summary from the reviewed findings so it matches what follows
            context = {"model_name": run.model_name}
            doc.add_paragraph(TemplateNarrativeWriter().write(r.agent, context, reviewed, {}).summary)
        else:
            doc.add_paragraph(r.summary)
            _para(doc, "AI summary written before human review; the reviewed status of each finding is shown below.",
                  italic=True, size=8, color=INK_2)
        for heading, img, caption in chart_list:
            if heading == r.agent:
                doc.add_picture(str(img), width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _para(doc, caption, italic=True, size=8, color=INK_2)
        for x in [x for x in counted if x.finding.agent == r.agent]:
            f, rv = x.finding, x.review
            doc.add_heading(f"{f.finding_id} · {f.title}", 3)
            meta = f"Severity: {f.severity.value}"
            if x.original_severity != f.severity:
                meta += f" (changed by reviewer from {x.original_severity.value})"
            meta += f" · {f.pillar.value} · Review: {rv.status.value}"
            if x.remediation_due:
                meta += f" · Remediation due {x.remediation_due}"
            _para(doc, meta, size=8.5, color=INK_2)
            doc.add_paragraph(f.description)
            ev = _evidence_lines(f.evidence)
            if ev:
                _para(doc, "Evidence", bold=True, size=9)
                for line in ev:
                    doc.add_paragraph(line, style="List Bullet").runs[0].font.size = Pt(8.5)
            if f.commentary:
                _labelled(doc, "AI commentary", f.commentary)
            _labelled(doc, "Recommendation", f.recommendation)
            if rv.comment:
                _labelled(doc, f"Validator note ({rv.reviewer})", rv.comment)

    # -- appendices --------------------------------------------------------------------------------------
    rejected = [x for x in finals_all if x.review.status is ReviewStatus.REJECTED]
    doc.add_heading("Appendix A. Findings rejected by the validator", 1)
    if rejected:
        rows = [[x.finding.finding_id, x.original_severity.value, x.finding.title, x.review.reviewer or "", x.review.comment or ""]
                for x in rejected]
        _table(doc, ["ID", "Proposed severity", "Finding", "Reviewer", "Rationale"], rows, [0.6, 0.8, 2.2, 1.0, 2.0])
    else:
        doc.add_paragraph("None.")

    doc.add_heading("Appendix B. Audit trail and turnaround", 1)
    ta = turnaround(run_dir)
    rows = [["Automated analysis", f"{ta.get('automated_seconds', '—')} s"]]
    for agent, secs in ta.get("agent_seconds", {}).items():
        rows.append([f"  {agent}", f"{secs} s"])
    rows += [
        ["Review decisions recorded", str(ta.get("decisions", 0))],
        ["Active review (first to last decision)", f"{ta['active_review_minutes']} min" if "active_review_minutes" in ta else "—"],
        ["Run completion to sign-off", f"{ta['review_hours']} h" if "review_hours" in ta else "Not signed off"],
        ["Total elapsed (start to sign-off)", f"{ta['total_hours']} h" if "total_hours" in ta else "Not signed off"],
    ]
    _table(doc, ["Measure", "Value"], rows, [3.6, 2.0])
    _para(doc, "Times are taken from the run's append-only audit log (audit.jsonl), which also records every "
          "review decision with reviewer, timestamp and rationale.", size=8.5, color=INK_2)
    if signed and signed.comment:
        _labelled(doc, "Sign-off comment", signed.comment)

    doc.add_heading("Appendix C. Limitations of this validation", 1)
    for text in (
        "Automated tests cover the checks listed in section 2; they do not replace the validator's judgement on "
        "conceptual soundness, business use or the appropriateness of the model's design.",
        "Statistical tests reject a correct model at roughly their significance level; findings are therefore "
        "reviewed individually before they are accepted.",
        "AI-written text (area summaries, commentary, executive summary) is generated from the computed results "
        "only and is labelled as such; it cannot add, remove or re-grade findings.",
    ):
        doc.add_paragraph(text, style="List Bullet")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def convert_to_pdf(docx_path: str | Path) -> Path | None:
    """Convert with LibreOffice if it is installed; returns None otherwise."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe is None:
        mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        exe = str(mac) if mac.exists() else None
    if exe is None:
        return None
    docx_path = Path(docx_path)
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                   check=True, capture_output=True, timeout=180)
    pdf = docx_path.with_suffix(".pdf")
    return pdf if pdf.exists() else None
